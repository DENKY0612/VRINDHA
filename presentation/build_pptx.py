#!/usr/bin/env python3
"""Build Vrindha AI competition presentation (16 slides)."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import nsmap
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

_HERE = Path(__file__).resolve().parent
ASSETS = _HERE / "assets"
OUT = _HERE / "Vrindha_AI_Presentation.pptx"

W, H = Inches(13.333), Inches(7.5)
NAVY = RGBColor(0x08, 0x12, 0x22)
NAVY2 = RGBColor(0x0D, 0x1B, 0x30)
CARD = RGBColor(0x12, 0x24, 0x3C)
CYAN = RGBColor(0x22, 0xD3, 0xEE)
TEAL = RGBColor(0x2D, 0xD4, 0xBF)
GOLD = RGBColor(0xE8, 0xC3, 0x6A)
WHITE = RGBColor(0xF5, 0xF7, 0xFA)
MUTED = RGBColor(0x9A, 0xAB, 0xC4)
SOFT = RGBColor(0xCB, 0xD5, 0xE1)
LINE = RGBColor(0x1E, 0x3A, 0x5F)


def _set_run(run, text, size, color, bold=False, italic=False, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font


def add_text(slide, l, t, w, h, text, size=18, color=WHITE, bold=False, italic=False,
             align=PP_ALIGN.LEFT, font="Calibri", anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    _set_run(p.add_run(), text, size, color, bold, italic, font)
    return box


def add_para(tf, text, size=16, color=WHITE, bold=False, italic=False, align=PP_ALIGN.LEFT, space_before=0, space_after=6):
    p = tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    _set_run(p.add_run(), text, size, color, bold, italic)
    return p


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def rect(slide, l, t, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    fill(s, color)
    return s


def rrect(slide, l, t, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    fill(s, color)
    s.adjustments[0] = 0.08
    return s


def line(slide, l, t, w, color=CYAN, thickness=Pt(3)):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, thickness)
    fill(s, color)
    return s


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, NAVY)
    return s


def chrome(slide, num, total=16):
    rect(slide, 0, 0, Inches(0.08), H, CYAN)
    add_text(slide, Inches(0.45), Inches(7.12), Inches(8), Inches(0.28),
             "VRINDHA AI  ·  Intelligent Cybersecurity & SOC System", 11, MUTED)
    add_text(slide, Inches(11.6), Inches(7.12), Inches(1.4), Inches(0.28),
             f"{num:02d}  /  {total:02d}", 11, MUTED, align=PP_ALIGN.RIGHT)


def section_label(slide, text, top=0.28):
    add_text(slide, Inches(0.5), Inches(top), Inches(12), Inches(0.28),
             text.upper(), 12, CYAN, bold=True)


def title(slide, text, top=0.52, size=30):
    add_text(slide, Inches(0.5), Inches(top), Inches(12.3), Inches(0.7),
             text, size, WHITE, bold=True)
    line(slide, Inches(0.5), Inches(top + 0.68), Inches(1.4), CYAN, Pt(3.5))


def card(slide, l, t, w, h, heading, body, accent=CYAN, tag=None):
    rrect(slide, l, t, w, h, CARD)
    rect(slide, l, t, Inches(0.07), h, accent)
    y = t + Inches(0.16)
    if tag:
        add_text(slide, l + Inches(0.22), y, w - Inches(0.35), Inches(0.24), tag, 11, accent, bold=True)
        y += Inches(0.26)
    add_text(slide, l + Inches(0.22), y, w - Inches(0.35), Inches(0.36), heading, 16, WHITE, bold=True)
    add_text(slide, l + Inches(0.22), y + Inches(0.36), w - Inches(0.35), h - (y - t) - Inches(0.48),
             body, 13, SOFT)


def pill(slide, l, t, w, h, text, bg=CYAN, fg=NAVY):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    fill(s, bg)
    s.adjustments[0] = 0.5
    tf = s.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    try:
        tf._txBody.bodyPr.set("anchor", "ctr")
    except Exception:
        pass
    _set_run(p.add_run(), text, 12, fg, True)
    return s


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # ── 1 TITLE ──────────────────────────────────────────────
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(str(ASSETS / "title_hero.png"), 0, 0, W, H)
    rect(s, 0, 0, W, H, RGBColor(0x04, 0x0A, 0x14))
    # dark overlay via semi-opaque feel: a left panel
    left = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(7.6), H)
    fill(left, NAVY)
    # restore image on right
    s.shapes.add_picture(str(ASSETS / "title_hero.png"), Inches(6.4), 0, Inches(6.94), H)
    rect(s, Inches(6.2), 0, Inches(0.35), H, NAVY)
    add_text(s, Inches(0.55), Inches(1.15), Inches(6.6), Inches(0.3),
             "COMPETITION PRESENTATION", 13, CYAN, bold=True)
    add_text(s, Inches(0.55), Inches(1.55), Inches(6.6), Inches(1.5),
             "Vrindha AI", 54, WHITE, bold=True)
    add_text(s, Inches(0.55), Inches(3.05), Inches(6.4), Inches(1.15),
             "Intelligent AI-Powered\nCybersecurity and SOC System", 22, GOLD, bold=False)
    line(s, Inches(0.55), Inches(4.35), Inches(1.6), CYAN, Pt(3.5))
    add_text(s, Inches(0.55), Inches(4.55), Inches(6.3), Inches(1.5),
             "What if a cybersecurity system could continuously monitor threats, understand what is happening, connect information, and recommend the right defensive action—while still keeping humans in control?",
             15, SOFT, italic=True)
    add_text(s, Inches(0.55), Inches(6.55), Inches(6.3), Inches(0.35),
             "Defensive  ·  Coordinated  ·  Human-authorized", 13, MUTED)
    notes(s, "What if a cybersecurity system could continuously monitor threats, understand what is happening, connect information from different systems, and recommend the right defensive action—while still keeping humans in control? Honourable judges, today we present Vrindha AI: an intelligent, AI-powered cybersecurity and SOC system.")

    # ── 2 GREETING ───────────────────────────────────────────
    s = blank(prs)
    chrome(s, 2)
    section_label(s, "01  ·  Introduction")
    title(s, "Greeting the bench")
    add_text(s, Inches(0.5), Inches(1.5), Inches(12.2), Inches(1.1),
             "Honourable judges, respected faculty, and fellow participants — thank you.\nWe are here to present Vrindha AI.", 20, WHITE)
    cards = [
        ("What we built", "An AI-assisted cybersecurity organization that helps a Security Operations Center detect, investigate, and defend."),
        ("How it thinks", "Not one isolated detector. A coordinated team of specialized AI modules that share intelligence."),
        ("Who decides", "Vrindha can analyze and recommend. High-impact actions stay under human authorization."),
    ]
    for i, (h, b) in enumerate(cards):
        card(s, Inches(0.5 + i * 4.2), Inches(3.0), Inches(3.95), Inches(3.15), h, b, tag=f"0{i+1}")
    notes(s, "Honourable judges, we present Vrindha AI — an intelligent cybersecurity platform designed to assist a Security Operations Center, threat intelligence operations, incident investigation, monitoring, risk analysis, and defensive response. Our long-term vision is a coordinated cybersecurity organization where specialized AI modules work together, while humans remain responsible for important and high-impact decisions.")

    # ── 3 PROBLEM ────────────────────────────────────────────
    s = blank(prs)
    chrome(s, 3)
    section_label(s, "02  ·  The problem")
    title(s, "Cybersecurity is no longer a single-alert problem")
    s.shapes.add_picture(str(ASSETS / "soc_problem.png"), Inches(8.15), Inches(1.55), Inches(4.7), Inches(5.15))
    points = [
        ("Alert volume is exploding", "Teams face more events than humans can review in time."),
        ("Threats move across systems", "An IP, a login, a file, and a network spike may be one attack."),
        ("Intelligence arrives too late", "Known threat data often sits outside the SOC workflow."),
        ("Tools do not talk to each other", "SIEM, TI feeds, and response playbooks stay isolated."),
    ]
    for i, (h, b) in enumerate(points):
        y = Inches(1.55 + i * 1.25)
        rrect(s, Inches(0.5), y, Inches(7.35), Inches(1.12), CARD)
        rect(s, Inches(0.5), y, Inches(0.08), Inches(1.12), GOLD if i % 2 else CYAN)
        add_text(s, Inches(0.8), y + Inches(0.16), Inches(6.8), Inches(0.35), h, 16, WHITE, bold=True)
        add_text(s, Inches(0.8), y + Inches(0.52), Inches(6.8), Inches(0.45), b, 14, SOFT)
    notes(s, "Cybersecurity today is not a shortage of tools. It is a shortage of coordination. Attacks move faster than isolated dashboards. A SOC — a Security Operations Center — is the team that watches alerts and investigates incidents. At scale, that work becomes overwhelming.")

    # ── 4 TRADITIONAL SOC ────────────────────────────────────
    s = blank(prs)
    chrome(s, 4)
    section_label(s, "03  ·  Why scale breaks traditional SOCs")
    title(s, "Traditional SOC operations struggle at scale")
    add_text(s, Inches(0.5), Inches(1.45), Inches(12.2), Inches(0.45),
             "A SOC is the security control room. At scale, the room fills faster than people can decide.", 16, SOFT)
    items = [
        ("01", "Too many alerts", "Analysts spend hours triaging noise instead of investigating real incidents."),
        ("02", "Slow correlation", "Connecting an alert to a known threat still depends on manual lookup."),
        ("03", "Isolated tools", "SIEM, threat feeds, and response systems do not share a common intelligence layer."),
        ("04", "Knowledge walks away", "Lessons from past incidents stay in tickets, chats, or one analyst’s memory."),
        ("05", "Fatigue and delay", "The more the environment grows, the longer it takes to understand what matters."),
        ("06", "Weak feedback loop", "Yesterday’s incident rarely improves tomorrow’s detection automatically."),
    ]
    for i, (n, h, b) in enumerate(items):
        col, row = i % 3, i // 3
        card(s, Inches(0.5 + col * 4.2), Inches(2.15 + row * 2.3), Inches(4.0), Inches(2.1), h, b, tag=n)
    notes(s, "Why do traditional SOC operations become difficult at scale? Because the work is still mostly sequential and human-limited: see an alert, search another system, check a threat feed, write notes, then decide. When telemetry multiplies, that chain breaks. We needed a coordinated intelligence layer — not another isolated detector.")

    # ── 5 INTRODUCING ────────────────────────────────────────
    s = blank(prs)
    chrome(s, 5)
    section_label(s, "04  ·  The project")
    title(s, "Introducing Vrindha AI")
    add_text(s, Inches(0.5), Inches(1.5), Inches(12.2), Inches(0.85),
             "Vrindha is a defensive cybersecurity platform that helps a SOC investigate, correlate, and respond — with specialized AI roles working as one organization.",
             18, WHITE)
    pillars = [
        ("Detect", "Watch telemetry, logs, and infrastructure for suspicious activity."),
        ("Understand", "Check whether an indicator is known and what it means."),
        ("Coordinate", "Share context across SOC, threat intelligence, and risk."),
        ("Recommend", "Prepare a defensive response for a human to approve."),
    ]
    for i, (h, b) in enumerate(pillars):
        rrect(s, Inches(0.5 + i * 3.15), Inches(2.6), Inches(3.0), Inches(2.55), CARD)
        circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.72 + i * 3.15), Inches(2.82), Inches(0.42), Inches(0.42))
        fill(circ, CYAN)
        tf = circ.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        try:
            tf._txBody.bodyPr.set("anchor", "ctr")
        except Exception:
            pass
        _set_run(p.add_run(), str(i + 1), 14, NAVY, True)
        add_text(s, Inches(1.25 + i * 3.15), Inches(2.86), Inches(2.05), Inches(0.38), h, 18, WHITE, bold=True)
        add_text(s, Inches(0.7 + i * 3.15), Inches(3.45), Inches(2.6), Inches(1.45), b, 14, SOFT)
    add_text(s, Inches(0.5), Inches(5.45), Inches(12.2), Inches(1.2),
             "The key innovation of our project is not “AI detects attacks.”\nIt is the coordination of multiple specialized AI capabilities — with accountability.",
             18, GOLD, italic=True)
    notes(s, "This is Vrindha AI. Instead of treating cybersecurity tools as isolated systems, we are creating a coordinated intelligence layer. Vrindha assists SOC operations, threat intelligence, investigation, risk analysis, and defensive response. It is designed to amplify cybersecurity professionals — not replace them.")

    # ── 6 HOW IT WORKS ───────────────────────────────────────
    s = blank(prs)
    chrome(s, 6)
    section_label(s, "05  ·  Operating model")
    title(s, "How Vrindha works")
    add_text(s, Inches(0.5), Inches(1.45), Inches(12.2), Inches(0.4),
             "A simple loop. Every step feeds the next.", 16, SOFT)
    steps = [
        ("1", "Detect", "SOC sees suspicious activity"),
        ("2", "Understand", "Threat Intelligence checks the indicator"),
        ("3", "Correlate", "Events from many systems are linked"),
        ("4", "Analyze", "Risk engine scores severity"),
        ("5", "Recommend", "SOC Analyst AI investigates"),
        ("6", "Approve", "Human authorizes high-impact action"),
        ("7", "Defend", "Commander AI coordinates response"),
        ("8", "Learn", "Knowledge AI stores lessons"),
    ]
    for i, (n, h, b) in enumerate(steps):
        col, row = i % 4, i // 4
        x, y = Inches(0.45 + col * 3.2), Inches(2.05 + row * 2.25)
        rrect(s, x, y, Inches(3.05), Inches(2.0), CARD)
        circ = s.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.18), y + Inches(0.22), Inches(0.42), Inches(0.42))
        fill(circ, GOLD if n == "6" else CYAN)
        tf = circ.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        try:
            tf._txBody.bodyPr.set("anchor", "ctr")
        except Exception:
            pass
        _set_run(p.add_run(), n, 14, NAVY, True)
        add_text(s, x + Inches(0.72), y + Inches(0.26), Inches(2.15), Inches(0.38), h, 18, WHITE, bold=True)
        add_text(s, x + Inches(0.2), y + Inches(0.85), Inches(2.65), Inches(0.9), b, 14, SOFT)
    notes(s, "Security telemetry arrives. The SOC detects suspicious activity. Threat Intelligence checks whether the indicator is known. The risk engine evaluates severity. The SOC Analyst AI investigates. Commander AI coordinates the response. A human administrator approves high-impact actions. Then the system stores the incident and learns. Detect. Understand. Correlate. Analyze. Recommend. Human approval. Defend. Learn.")

    # ── 7 SOC + TI ───────────────────────────────────────────
    s = blank(prs)
    chrome(s, 7)
    section_label(s, "06  ·  Integration")
    title(s, "SOC + Threat Intelligence, working as one")
    add_text(s, Inches(0.5), Inches(1.45), Inches(12.2), Inches(0.7),
             "Threat Intelligence is the process of collecting and understanding information about cyber threats so the organization can recognize attacks faster.",
             16, SOFT)

    rrect(s, Inches(0.5), Inches(2.3), Inches(5.7), Inches(4.2), CARD)
    add_text(s, Inches(0.75), Inches(2.5), Inches(5.2), Inches(0.35), "SOC  ·  Security Operations", 14, CYAN, bold=True)
    add_text(s, Inches(0.75), Inches(2.95), Inches(5.2), Inches(0.4), "Watches the environment", 20, WHITE, bold=True)
    for i, t in enumerate([
        "Collects logs, alerts, and telemetry",
        "Flags suspicious activity in real time",
        "Correlates security events",
        "Assists human analysts during incidents",
    ]):
        add_text(s, Inches(0.85), Inches(3.5 + i * 0.6), Inches(5.1), Inches(0.5), "▸   " + t, 15, SOFT)

    rrect(s, Inches(6.5), Inches(2.3), Inches(6.3), Inches(4.2), CARD)
    add_text(s, Inches(6.75), Inches(2.5), Inches(5.8), Inches(0.35), "TI  ·  Threat Intelligence", 14, GOLD, bold=True)
    add_text(s, Inches(6.75), Inches(2.95), Inches(5.8), Inches(0.4), "Explains what the signal means", 20, WHITE, bold=True)
    for i, t in enumerate([
        "Identifies indicators of compromise (IOCs)",
        "Checks if an IP, domain, or file is known-bad",
        "Tracks emerging campaigns and actors",
        "Shares useful intelligence back to the SOC",
    ]):
        add_text(s, Inches(6.85), Inches(3.5 + i * 0.6), Inches(5.7), Inches(0.5), "▸   " + t, 15, SOFT)
    notes(s, "A SOC watches what is happening inside the organization. Threat Intelligence explains whether that activity is known in the wider threat landscape. An indicator of compromise — an IOC — is a clue such as a malicious IP or file hash. When these two work together, a suspicious IP is no longer just an alert. It becomes context: known actor, known campaign, known risk.")

    # ── 8 MODULES ────────────────────────────────────────────
    s = blank(prs)
    chrome(s, 8)
    section_label(s, "07  ·  Architecture")
    title(s, "Specialized AI modules — one organization")
    mods = [
        ("Commander AI", "Oversees the system, coordinates information, and directs defensive response."),
        ("Threat Intelligence AI", "Collects threat data, identifies IOCs, and shares useful intelligence."),
        ("SOC Analyst AI", "Investigates alerts, correlates events, and recommends response steps."),
        ("Data Science AI", "Analyzes security data and improves anomaly detection and risk models."),
        ("Infrastructure AI", "Monitors servers, databases, networks, and service health."),
        ("Knowledge AI", "Stores procedures, past incidents, and lessons learned."),
        ("Ethics & Compliance AI", "Checks authorization, policy, and ethical limits before action."),
    ]
    # 4 on first row, 3 on second
    for i, (h, b) in enumerate(mods[:4]):
        card(s, Inches(0.4 + i * 3.22), Inches(1.5), Inches(3.08), Inches(2.4), h, b, accent=CYAN)
    for i, (h, b) in enumerate(mods[4:]):
        card(s, Inches(1.95 + i * 3.22), Inches(4.1), Inches(3.08), Inches(2.4), h, b, accent=GOLD if "Ethics" in h else TEAL)
    notes(s, "Vrindha is not one model doing everything. Commander AI oversees. Threat Intelligence AI collects and explains threats. SOC Analyst AI investigates. Data Science AI improves detection. Infrastructure AI watches systems. Knowledge AI remembers incidents. Ethics and Compliance AI prevents inappropriate autonomous decisions. Each has a role. Together they form a team.")

    # ── 9 HIVE ───────────────────────────────────────────────
    s = blank(prs)
    chrome(s, 9)
    section_label(s, "08  ·  Central intelligence")
    title(s, "Hive intelligence — a coordinated team")
    s.shapes.add_picture(str(ASSETS / "hive_network.png"), Inches(7.85), Inches(1.55), Inches(5.05), Inches(5.15))
    add_text(s, Inches(0.5), Inches(1.5), Inches(7.15), Inches(0.85),
             "Each AI has a specific responsibility. They communicate. The center correlates.",
             16, SOFT)
    hive = [
        "Correlate events from multiple systems",
        "Detect emerging attack patterns",
        "Share defensive intelligence",
        "Improve models using historical incidents",
        "Generate situation reports for administrators",
        "Recommend response strategies to humans",
    ]
    for i, t in enumerate(hive):
        y = Inches(2.4 + i * 0.68)
        rrect(s, Inches(0.5), y, Inches(7.15), Inches(0.6), CARD)
        circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.68), y + Inches(0.16), Inches(0.28), Inches(0.28))
        fill(circ, CYAN)
        add_text(s, Inches(1.15), y + Inches(0.12), Inches(6.3), Inches(0.4), t, 15, WHITE)
    notes(s, "The system works like hive intelligence — a coordinated team of specialized agents. Information from one module improves another. A suspicious IP found by the SOC is sent to Threat Intelligence. If it is a known threat, the risk score rises. The Analyst AI investigates related events. Commander AI coordinates. Knowledge AI stores the lesson. That is how the hive learns.")

    # ── 10 DISTRIBUTED ───────────────────────────────────────
    s = blank(prs)
    chrome(s, 10)
    section_label(s, "09  ·  Future network")
    title(s, "A distributed, authorized cybersecurity network")
    add_text(s, Inches(0.5), Inches(1.45), Inches(12.2), Inches(0.55),
             "Authorized computers can join the organization. This is not designed for unauthorized access.",
             16, SOFT)
    contrib = [
        ("Telemetry", "Security signals from member systems"),
        ("Logs & alerts", "What each node is seeing"),
        ("Threat intel", "Indicators and campaign data"),
        ("Detection rules", "Shared defensive knowledge"),
        ("Model updates", "Safer improvements over time"),
        ("Incidents", "Lessons that raise the whole hive"),
    ]
    for i, (h, b) in enumerate(contrib):
        col, row = i % 3, i // 3
        card(s, Inches(0.5 + col * 4.2), Inches(2.15 + row * 1.7), Inches(4.0), Inches(1.55), h, b)
    add_text(s, Inches(0.5), Inches(5.7), Inches(12.2), Inches(0.9),
             "Participation is voluntary  ·  authorized  ·  authenticated  ·  encrypted  ·  audited",
             16, GOLD, bold=True)
    notes(s, "The future vision is a network of authorized computers. Each member can contribute telemetry, logs, threat intelligence, malware indicators, detection rules, model updates, and incident information. The central AI correlates that information and sends useful defensive intelligence back. Participation is voluntary, authorized, authenticated, encrypted, and audited. This is not unauthorized access to computers.")

    # ── 11 ETHICS ────────────────────────────────────────────
    s = blank(prs)
    chrome(s, 11)
    section_label(s, "10  ·  Safety principle")
    title(s, "Ethics, authorization, and human approval")
    s.shapes.add_picture(str(ASSETS / "human_control.png"), Inches(8.2), Inches(1.55), Inches(4.65), Inches(5.15))
    add_text(s, Inches(0.5), Inches(1.5), Inches(7.4), Inches(1.15),
             "Vrindha can detect, analyze, investigate, recommend, and prepare a response. Important actions remain under human authorization.",
             17, WHITE)
    rules = [
        ("Defensive by design", "Vrindha is not built to autonomously launch attacks or run unauthorized offensive operations."),
        ("High-impact needs a human", "Blocking, containment, and any action with real consequence requires approval."),
        ("Intelligence without control is dangerous", "That is why human authorization is a core part of our architecture."),
    ]
    for i, (h, b) in enumerate(rules):
        y = Inches(2.85 + i * 1.2)
        rrect(s, Inches(0.5), y, Inches(7.45), Inches(1.08), CARD)
        rect(s, Inches(0.5), y, Inches(0.08), Inches(1.08), GOLD)
        add_text(s, Inches(0.8), y + Inches(0.14), Inches(6.95), Inches(0.32), h, 15, GOLD, bold=True)
        add_text(s, Inches(0.8), y + Inches(0.48), Inches(6.95), Inches(0.5), b, 13, SOFT)
    notes(s, "The most important safety principle: Vrindha is designed primarily for defensive cybersecurity. It should not autonomously launch attacks or perform unauthorized offensive operations. For high-impact actions, human approval remains required. We are not building an uncontrolled autonomous hacker. We are building an AI-assisted cybersecurity organization with accountability and safety controls.")

    # ── 12 INNOVATION ────────────────────────────────────────
    s = blank(prs)
    chrome(s, 12)
    section_label(s, "11  ·  Differentiation")
    title(s, "What makes Vrindha innovative")
    add_text(s, Inches(0.5), Inches(1.45), Inches(12.2), Inches(0.45),
             "The innovation is not “AI detects attacks.” The innovation is coordination.", 17, GOLD, italic=True)

    left_items = [
        ("Isolated tools", "Each product sees only its own slice of the problem."),
        ("One-shot detection", "A model fires an alert and the story ends there."),
        ("Manual stitching", "Humans copy context from one screen to another."),
        ("No shared memory", "Incidents do not reliably improve the next detection."),
    ]
    right_items = [
        ("Coordinated intelligence", "SOC, TI, risk, and response share one loop."),
        ("Detect → Learn cycle", "Every incident can raise future detection quality."),
        ("Specialized roles", "Each AI does one job well, then reports to the hive."),
        ("Human remains accountable", "Recommendation is automated. Authority is not."),
    ]
    add_text(s, Inches(0.55), Inches(2.05), Inches(5.8), Inches(0.35), "COMMON APPROACH", 12, MUTED, bold=True)
    add_text(s, Inches(6.95), Inches(2.05), Inches(5.8), Inches(0.35), "VRINDHA APPROACH", 12, CYAN, bold=True)
    for i, ((h1, b1), (h2, b2)) in enumerate(zip(left_items, right_items)):
        y = Inches(2.45 + i * 1.05)
        rrect(s, Inches(0.5), y, Inches(5.85), Inches(0.95), CARD)
        add_text(s, Inches(0.7), y + Inches(0.1), Inches(5.5), Inches(0.3), h1, 14, WHITE, bold=True)
        add_text(s, Inches(0.7), y + Inches(0.42), Inches(5.5), Inches(0.42), b1, 13, MUTED)
        rrect(s, Inches(6.9), y, Inches(5.9), Inches(0.95), CARD)
        rect(s, Inches(6.9), y, Inches(0.08), Inches(0.95), CYAN)
        add_text(s, Inches(7.15), y + Inches(0.1), Inches(5.45), Inches(0.3), h2, 14, WHITE, bold=True)
        add_text(s, Inches(7.15), y + Inches(0.42), Inches(5.45), Inches(0.42), b2, 13, SOFT)
    notes(s, "What makes Vrindha different is the coordination of specialized AI capabilities. A suspicious IP detected by the SOC can be sent to Threat Intelligence. Threat Intelligence can determine whether that IP is associated with a known threat. The risk engine can then raise or lower the score. The SOC Analyst AI investigates related events. Commander AI coordinates. The system stores the incident and uses the lesson to improve future detection.")

    # ── 13 FUTURE ────────────────────────────────────────────
    s = blank(prs)
    chrome(s, 13)
    section_label(s, "12  ·  Roadmap")
    title(s, "Future scope")
    future = [
        ("01", "Richer hive", "More specialized modules joining the same authorized event bus."),
        ("02", "Better models", "Data Science AI continuously improving anomaly and risk scoring from real incidents."),
        ("03", "Authorized nodes", "A growing network of voluntary, authenticated member systems."),
        ("04", "Shared defense", "Detection rules and indicators flowing back to members safely."),
        ("05", "Situation reports", "Clear briefings for administrators, not only raw alerts."),
        ("06", "Stronger audit", "Every recommendation and approval remains explainable and logged."),
    ]
    for i, (n, h, b) in enumerate(future):
        col, row = i % 3, i // 3
        card(s, Inches(0.5 + col * 4.2), Inches(1.55 + row * 2.5), Inches(4.0), Inches(2.3), h, b, tag=n)
    notes(s, "Looking ahead, Vrindha can grow as a coordinated organization: more authorized nodes, stronger models trained on historical incidents, richer situation reports, and safer distribution of defensive intelligence. The architecture is designed so new modules can join without bypassing safety, authorization, or human approval.")

    # ── 14 IMPACT ────────────────────────────────────────────
    s = blank(prs)
    chrome(s, 14)
    section_label(s, "13  ·  Why this matters")
    title(s, "Real-world impact")
    impacts = [
        ("Faster understanding", "Analysts spend less time stitching tools and more time deciding."),
        ("Fewer missed connections", "An alert is checked against known threats before it is dismissed."),
        ("Safer response", "Recommendations are prepared, but high-impact action stays authorized."),
        ("Institutional memory", "Incidents become lessons the system can reuse."),
        ("Lower fatigue", "Coordination reduces noise and highlights what is actually urgent."),
        ("Human amplification", "Our goal is not to replace professionals. It is to amplify them."),
    ]
    for i, (h, b) in enumerate(impacts):
        col, row = i % 3, i // 3
        card(s, Inches(0.5 + col * 4.2), Inches(1.55 + row * 2.5), Inches(4.0), Inches(2.3), h, b, accent=TEAL)
    notes(s, "In the real world, this means a SOC that understands faster, misses fewer connections, responds more safely, and remembers what it has already learned. We are not claiming a system that stops every attack. We are building a system that makes defensive teams stronger, calmer, and more accountable.")

    # ── 15 CONCLUSION ────────────────────────────────────────
    s = blank(prs)
    chrome(s, 15)
    section_label(s, "14  ·  Closing")
    title(s, "The future is not humans versus AI")
    s.shapes.add_picture(str(ASSETS / "closing_together.png"), Inches(8.15), Inches(1.55), Inches(4.7), Inches(5.15))
    add_text(s, Inches(0.5), Inches(1.6), Inches(7.4), Inches(2.3),
             "Our vision with Vrindha is not to build an AI that acts without limits. Our vision is to build an AI that understands threats, coordinates intelligence, assists humans, learns from incidents, and protects systems responsibly.",
             18, WHITE)
    rrect(s, Inches(0.5), Inches(4.15), Inches(7.4), Inches(2.35), CARD)
    rect(s, Inches(0.5), Inches(4.15), Inches(0.08), Inches(2.35), GOLD)
    add_text(s, Inches(0.8), Inches(4.35), Inches(6.9), Inches(1.95),
             "Because the future of cybersecurity is not humans versus AI.\n\nIt is humans and AI working together — with intelligence, authorization, and responsibility.",
             17, GOLD, italic=True)
    notes(s, "Our vision with Vrindha is not to build an AI that acts without limits. Our vision is to build an AI that understands threats, coordinates intelligence, assists humans, learns from incidents, and protects systems responsibly. Because the future of cybersecurity is not humans versus AI. It is humans and AI working together — with intelligence, authorization, and responsibility.")

    # ── 16 THANK YOU ─────────────────────────────────────────
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(str(ASSETS / "title_hero.png"), 0, 0, W, H)
    overlay = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    fill(overlay, NAVY)
    # keep a band of the hero
    s.shapes.add_picture(str(ASSETS / "title_hero.png"), Inches(8.4), 0, Inches(4.94), H)
    rect(s, Inches(8.15), 0, Inches(0.4), H, NAVY)
    add_text(s, Inches(0.7), Inches(1.7), Inches(7.2), Inches(0.35), "THANK YOU", 14, CYAN, bold=True)
    add_text(s, Inches(0.7), Inches(2.15), Inches(7.2), Inches(1.1), "Questions welcome.", 44, WHITE, bold=True)
    line(s, Inches(0.7), Inches(3.4), Inches(1.5), GOLD, Pt(3.5))
    add_text(s, Inches(0.7), Inches(3.7), Inches(7.1), Inches(1.4),
             "Vrindha AI — defensive cybersecurity with coordinated intelligence and human authorization.",
             18, SOFT)
    add_text(s, Inches(0.7), Inches(5.3), Inches(7.1), Inches(0.9),
             "Detect  →  Understand  →  Correlate  →  Analyze\nRecommend  →  Human Approval  →  Defend  →  Learn",
             16, GOLD)
    add_text(s, Inches(0.7), Inches(6.55), Inches(7.1), Inches(0.35),
             "Thank you, judges and audience.", 14, MUTED)
    notes(s, "Thank you, honourable judges and the audience. We welcome your questions.")

    prs.save(str(OUT))
    print(f"Saved {OUT}")


if __name__ == "__main__":
    build()
